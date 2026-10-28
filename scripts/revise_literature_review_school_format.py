from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path("/Users/verneywang/Downloads/研究生学位（毕业）论文文献综述_修订版.docx")
OUTPUT = Path("/Users/verneywang/Downloads/研究生学位（毕业）论文文献综述_学校格式修订版.docx")


# The original bibliography is numbered in list order, but the body first cites
# references 27–31 before 23–26. Reference 33 (AgentTuning) is outside the
# review's stated non-fine-tuning scope and is never cited, so it is removed.
# The remaining items are renumbered by first appearance.
OLD_TO_NEW = {
    **{i: i for i in range(1, 23)},
    27: 23,
    28: 24,
    29: 25,
    30: 26,
    31: 27,
    23: 28,
    24: 29,
    25: 30,
    26: 31,
    36: 32,
    32: 33,
    34: 34,
    35: 35,
}


REFERENCES = [
    "WANG L, MA C, FENG X, et al. A survey on large language model based autonomous agents[J]. Frontiers of Computer Science, 2024, 18(6): 186345.",
    "XI Z, CHEN W, GUO X, et al. The rise and potential of large language model based agents: A survey[J]. Science China Information Sciences, 2025, 68(2): 121101.",
    "YAO S, ZHAO J, YU D, et al. ReAct: Synergizing reasoning and acting in language models[A]. Proceedings of the Eleventh International Conference on Learning Representations[C]. Kigali: OpenReview, 2023.",
    "GAO L, MADAAN A, ZHOU S, et al. PAL: Program-aided language models[A]. Proceedings of the 40th International Conference on Machine Learning[C]. Honolulu: PMLR, 2023, 202: 10764-10799.",
    "KHATTAB O, SINGHVI A, MAHESHWARI P, et al. DSPy: Compiling declarative language model calls into self-improving pipelines[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "HU S, LU C, CLUNE J. Automated design of agentic systems[A]. Proceedings of the Thirteenth International Conference on Learning Representations[C]. Singapore: OpenReview, 2025.",
    "ZHANG J, XIANG J, YU Z, et al. AFlow: Automating agentic workflow generation[A]. Proceedings of the Thirteenth International Conference on Learning Representations[C]. Singapore: OpenReview, 2025.",
    "WANG G, XIE Y, JIANG Y, et al. Voyager: An open-ended embodied agent with large language models[J]. Transactions on Machine Learning Research, 2024.",
    "ZHAO M, WEI X, SHAO Y, et al. A²Flow: Automating agentic workflow generation via self-adaptive abstraction operators[A]. Proceedings of the Fortieth AAAI Conference on Artificial Intelligence[C]. Palo Alto: AAAI Press, 2026: 29930-29938.",
    "SCHICK T, DWIVEDI-YU J, DESSÌ R, et al. Toolformer: Language models can teach themselves to use tools[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 68539-68551.",
    "ZHENG Y, ZHANG Z, MA C, et al. SkillRouter: Skill routing for LLM agents at scale[EB/OL]. arXiv:2603.22455, 2026-07-20[2026-08-25]. https://arxiv.org/abs/2603.22455.",
    "WEI J, WANG X, SCHUURMANS D, et al. Chain-of-thought prompting elicits reasoning in large language models[A]. Advances in Neural Information Processing Systems 35[C]. Red Hook: Curran Associates, 2022: 24824-24837.",
    "WANG X, WEI J, SCHUURMANS D, et al. Self-consistency improves chain of thought reasoning in language models[A]. Proceedings of the Eleventh International Conference on Learning Representations[C]. Kigali: OpenReview, 2023.",
    "YAO S, YU D, ZHAO J, et al. Tree of thoughts: Deliberate problem solving with large language models[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 11809-11822.",
    "SHINN N, CASSANO F, GOPINATH A, et al. Reflexion: Language agents with verbal reinforcement learning[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 8634-8652.",
    "MADAAN A, TANDON N, GUPTA P, et al. Self-refine: Iterative refinement with self-feedback[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 46534-46594.",
    "PRYZANT R, ITER D, LI J, et al. Automatic prompt optimization with \"gradient descent\" and beam search[A]. BOUAMOR H, PINO J, BALI K, eds. Proceedings of the 2023 Conference on Empirical Methods in Natural Language Processing[C]. Singapore: Association for Computational Linguistics, 2023: 7957-7968.",
    "YANG C, WANG X, LU Y, et al. Large language models as optimizers[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "GUO Q, WANG R, GUO J, et al. EvoPrompt: Connecting LLMs with evolutionary algorithms yields powerful prompt optimizers[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "FERNANDO C, BANARSE D S, MICHALEWSKI H, et al. Promptbreeder: Self-referential self-improvement via prompt evolution[A]. SALAKHUTDINOV R, KOLTER Z, HELLER K, et al., eds. Proceedings of the 41st International Conference on Machine Learning[C]. Vienna: PMLR, 2024, 235: 13481-13544.",
    "YUKSEKGONUL M, BIANCHI F, BOEN J, et al. Optimizing generative AI by backpropagating language model feedback[J]. Nature, 2025, 639: 609-616.",
    "LI Z, XU S, MEI K, et al. AutoFlow: Automated workflow generation for large language model agents[EB/OL]. arXiv:2407.12821, 2024-07-01[2026-08-25]. https://arxiv.org/abs/2407.12821.",
    "SHEN Y, SONG K, TAN X, et al. HuggingGPT: Solving AI tasks with ChatGPT and its friends in Hugging Face[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 38154-38180.",
    "PARANJAPE B, LUND B, SINGH S, et al. ART: Automatic multi-step reasoning and tool-use for large language models[EB/OL]. arXiv:2303.09014, 2023-03-16[2026-08-25]. https://arxiv.org/abs/2303.09014.",
    "PATIL S G, ZHANG T, WANG X, et al. Gorilla: Large language model connected with massive APIs[A]. Advances in Neural Information Processing Systems 37[C]. Red Hook: Curran Associates, 2024: 126544-126565.",
    "QIN Y, HU S, LIN Y, et al. Tool learning with foundation models[J]. ACM Computing Surveys, 2024, 57(4): 1-40.",
    "QIN Y, LIANG S, YE Y, et al. ToolLLM: Facilitating large language models to master 16000+ real-world APIs[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "WU Q, BANSAL G, ZHANG J, et al. AutoGen: Enabling next-gen LLM applications via multi-agent conversation[A]. Proceedings of the First Conference on Language Modeling[C]. Philadelphia: OpenReview, 2024.",
    "LI G, HAMMOUD H A A K, ITANI H, et al. CAMEL: Communicative agents for \"mind\" exploration of large scale language model society[A]. Advances in Neural Information Processing Systems 36[C]. Red Hook: Curran Associates, 2023: 51991-52008.",
    "HONG S, ZHUGE M, CHEN J, et al. MetaGPT: Meta programming for a multi-agent collaborative framework[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "QIAN C, LIU W, LIU H, et al. ChatDev: Communicative agents for software development[A]. KU L W, MARTINS A, SRIKUMAR V, eds. Proceedings of the 62nd Annual Meeting of the Association for Computational Linguistics[C]. Bangkok: Association for Computational Linguistics, 2024: 15174-15186.",
    "GAO X. Compositional skill routing for LLM agents: Decompose, retrieve, and compose[EB/OL]. arXiv:2606.18051, 2026-06-16[2026-08-25]. https://arxiv.org/abs/2606.18051.",
    "LIU X, YU H, ZHANG H, et al. AgentBench: Evaluating LLMs as agents[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "MIALON G, FOURRIER C, SWIFT C, et al. GAIA: A benchmark for general AI assistants[A]. Proceedings of the Twelfth International Conference on Learning Representations[C]. Vienna: OpenReview, 2024.",
    "MA C, ZHANG J, ZHU Z, et al. AgentBoard: An analytical evaluation board of multi-turn LLM agents[A]. Advances in Neural Information Processing Systems 37[C]. Red Hook: Curran Associates, 2024: 74325-74362.",
]


def set_run_font(run, east_asia: str, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = east_asia
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), east_asia)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def expand_citation(spec: str) -> list[int]:
    nums: list[int] = []
    for part in re.split(r"[,，]", spec):
        if "-" in part:
            start, end = (int(x) for x in part.split("-", 1))
            nums.extend(range(start, end + 1))
        else:
            nums.append(int(part))
    return nums


def compress_citation(nums: list[int]) -> str:
    ordered = sorted(set(nums))
    pieces: list[str] = []
    start = prev = ordered[0]
    for n in ordered[1:] + [None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        pieces.append(str(start) if start == prev else f"{start}-{prev}")
        if n is not None:
            start = prev = n
    return ",".join(pieces)


def renumber_text(text: str) -> str:
    pattern = re.compile(r"\[(\d+(?:(?:-|,|，)\d+)*)\]")

    def repl(match: re.Match[str]) -> str:
        old_nums = expand_citation(match.group(1))
        if any(n not in OLD_TO_NEW for n in old_nums):
            raise ValueError(f"unmapped citation in body: {match.group(0)}")
        return f"[{compress_citation([OLD_TO_NEW[n] for n in old_nums])}]"

    return pattern.sub(repl, text)


def citation_numbers(text: str) -> list[int]:
    nums: list[int] = []
    for match in re.finditer(r"\[(\d+(?:(?:-|,|，)\d+)*)\]", text):
        nums.extend(expand_citation(match.group(1)))
    return nums


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if len(REFERENCES) != 35:
        raise AssertionError(f"expected 35 references, got {len(REFERENCES)}")

    doc = Document(SOURCE)
    heading_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "参考文献")

    # Clarify that reference 26 is a survey instead of silently grouping it with
    # four named systems. This makes every grouped citation semantically precise.
    for paragraph in doc.paragraphs[:heading_index]:
        for run in paragraph.runs:
            run.text = run.text.replace(
                "HuggingGPT、ART、Gorilla及ToolLLM[27-31]分别从模型协同、自动多步工具使用、API调用和大规模工具训练数据等角度扩展外部能力。",
                "HuggingGPT、ART、Gorilla、工具学习综述及ToolLLM[27-31]分别从模型协同、自动多步工具使用、API调用、工具学习综述和大规模工具训练数据等角度扩展外部能力。",
            )

    # Renumber citations without flattening the existing run-level formatting.
    for paragraph in doc.paragraphs[:heading_index]:
        for run in paragraph.runs:
            run.text = renumber_text(run.text)

    heading = doc.paragraphs[heading_index]
    if not heading.runs:
        heading.add_run("参考文献")
    for run in heading.runs:
        set_run_font(run, "黑体", 12, True)
    heading.paragraph_format.keep_with_next = True

    ref_paragraphs = [p for p in doc.paragraphs[heading_index + 1 :] if re.match(r"^\s*\[\d+\]", p.text)]
    if len(ref_paragraphs) != 36:
        raise AssertionError(f"expected 36 source reference paragraphs, got {len(ref_paragraphs)}")

    for idx, entry in enumerate(REFERENCES, start=1):
        paragraph = ref_paragraphs[idx - 1]
        paragraph.clear()
        run = paragraph.add_run(f"[{idx}] {entry}")
        set_run_font(run, "宋体", 10.5, False)
        fmt = paragraph.paragraph_format
        fmt.left_indent = Pt(21)
        fmt.first_line_indent = Pt(-21)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for paragraph in ref_paragraphs[len(REFERENCES) :]:
        delete_paragraph(paragraph)

    # Structural audit before saving.
    all_body_text = "\n".join(p.text for p in doc.paragraphs[:heading_index])
    used: list[int] = []
    first_seen: list[int] = []
    for n in citation_numbers(all_body_text):
        used.append(n)
        if n not in first_seen:
            first_seen.append(n)
    expected = list(range(1, 36))
    if first_seen != expected:
        raise AssertionError(f"first-appearance order is not sequential: {first_seen}")
    if sorted(set(used)) != expected:
        raise AssertionError(f"citation/reference coverage mismatch: {sorted(set(used))}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"references={len(REFERENCES)}; first_appearance={first_seen}")


if __name__ == "__main__":
    main()
