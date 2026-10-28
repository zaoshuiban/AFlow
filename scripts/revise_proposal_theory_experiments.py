from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import _Row
from docx.shared import Pt


SOURCE = Path("/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤.docx")
OUTPUT = Path("/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_研究方案与实验结果修订版.docx")


THEORY_ROWS = [
    (
        True,
        "2.1 模型构建的理论依据与形式化定义",
    ),
    (
        False,
        "（1）模块化程序合成与受约束搜索依据。智能体工作流可以抽象为由基础算子及其连接关系组成的可执行程序。ADAS与AFlow等研究表明，可以利用任务反馈在程序空间中搜索智能体结构[10,12]。本研究进一步把单个专项算子视为局部能力程序：稳定失败机制给出目标规格，参数化设计集提供输入输出样例，结构化算子模板限定接口、算法骨架、复杂度与安全边界，多候选生成和锦标赛则在受约束假设空间内近似搜索较优实现。与无约束自由代码生成相比，缩小假设空间有助于降低不可执行、超时和样例过拟合风险。",
    ),
    (
        False,
        "（2）独立验证与约束风险最小化依据。设基线工作流为B，目标机制为Mj，算子生成器仅根据机制说明和设计集Ddes生成候选集合Oj={oj1,…,ojK}，隐藏门控集Dgate和回归集Dreg不参与生成。候选质量以正确时间内完成率为核心，可表示为Q(o)=ΔAccgate+λ1·Stability−λ2·Regression−λ3·Timeout−λ4·Cost；晋级同时满足ΔAccgate≥τg、Regression≤τr、Stability≥τs且硬回归为0。该模型把“专项提升”写成带回归、时延和成本约束的经验风险最小化问题，而不是只追求开发集平均得分。",
    ),
    (
        False,
        "（3）带拒绝选项的选择性决策依据。设已晋级算子库为O={o1,…,oK}，路由器g(x)输出0表示使用基线，输出k表示尝试第k个算子。完整系统定义为：当g(x)=k、问题满足算子适用条件Ak(x)=1且执行成功时，F(x)=ok(x)；其余情况均令F(x)=B(x)。路由器只有在适用概率不低于τroute、预测风险不高于τrisk且输入未触及失败边界时才调用专项算子，否则主动拒绝并回退。这与带拒绝选项的选择性分类思想一致，使系统能够用覆盖率换取较低误触发和较小非专项风险。",
    ),
    (
        False,
        "（4）可证伪的系统净收益假设。研究不以“某个专项工具单独得分更高”作为最终结论，而检验H1：自动生成算子在独立隐藏机制题上优于原始AFlow；H2：证据感知路由能把专项增益转化为混合分布下的整体净收益；H3：算子数量增加时，误触发、跨算子冲突和非专项回归仍受预设阈值约束。由此，模型中的生成、晋级和路由三个环节均对应独立对照、消融指标和失败条件，研究方案可以被实验验证或否定。",
    ),
]


EXPERIMENT_ROWS = [
    (
        "实验协议。三组确认实验均使用候选生成器未见过的全新随机种子隐藏集，执行模型为deepseek-v4-flash，每道题独立运行3次，单次超时180秒，主指标为correct_within_timeout。实验前固定算子版本、路由规则、成功门槛和文件哈希；确认过程中不安装、不修改候选，使用math_equivalence_v7评分器，并保存逐题回答和路由记录。",
    ),
    (
        "（1）括号化表达式去重算子单算子确认（seed 120007）。在132次混合观测中，原始AFlow整体正确91/132（68.9%），接入路由后的系统正确117/132（88.6%）；专项题由8/36（22.2%）提高到36/36（100%），专项路由召回率为1.000，普通题与近似反例误路由为0，超时由29次降至3次，逐题胜/平/负为28/102/2。非专项题有2次新增错误，但均不是路由触发造成。",
    ),
    (
        "（2）受限数字素数枚举算子单算子确认（seed 150017）。在132次混合观测中，原始AFlow整体正确97/132（73.5%），路由系统正确129/132（97.7%）；专项题由4/36（11.1%）提高到36/36（100%），专项路由召回率为1.000，误路由为0，非专项新增错误为0，超时由32次降至0次，逐题胜/平/负为32/100/0。",
    ),
    (
        "（3）双算子联合确认（seed 160021）。在144次观测中，两类专项任务各24次，普通MATH题与两类近似反例共96次。原始AFlow整体正确96/144（66.7%），双算子路由系统正确132/144（91.7%）；括号化机制由9/24提高到24/24，受限数字素数机制由0/24提高到24/24，两类专项路由召回率均为1.000，普通题误路由和跨算子抢路由均为0，超时由39次降至1次，逐题胜/平/负为39/102/3。3次非专项新增错误经审计均不是路由触发造成。",
    ),
    (
        "阶段性结论。三组结果支持“已通过隐藏门控的专项算子能够接入AFlow并在参数化混合隐藏集上产生端到端净收益”，也初步说明两个算子共同存在时没有出现误路由或相互抢占。但这些结果仍属于MATH参数化问题族的可行性证据，不能替代完整MATH自然分布、更多机制、不同执行模型和跨领域实验；后续研究将按研究方案中的公平对照和冻结协议继续验证。",
    ),
]


def clear_and_set(paragraph, text: str, *, heading: bool) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.size = Pt(12)
    run.bold = heading
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = None if heading else Pt(24)


def insert_rows_after(table, anchor_row, template_row, items: list[tuple[bool, str]]) -> None:
    anchor_tr = anchor_row._tr
    for heading, text in items:
        new_tr = deepcopy(template_row._tr)
        anchor_tr.addnext(new_tr)
        new_row = _Row(new_tr, table)
        cell = new_row.cells[0]
        # The selected templates contain one paragraph. Keep cell properties and
        # row geometry, but replace the visible text and local paragraph format.
        paragraph = cell.paragraphs[0]
        clear_and_set(paragraph, text, heading=heading)
        for extra in cell.paragraphs[1:]:
            extra._element.getparent().remove(extra._element)
        anchor_tr = new_tr


def find_cell_paragraph(table, exact: str):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text.strip() == exact:
                    return paragraph
    raise ValueError(f"paragraph not found in table: {exact}")


def find_cell_paragraph_prefix(table, prefix: str):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text.strip().startswith(prefix):
                    return paragraph
    raise ValueError(f"paragraph prefix not found in table: {prefix}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)
    if len(doc.tables) != 8:
        raise AssertionError(f"unexpected table count: {len(doc.tables)}")

    scheme = doc.tables[2]
    foundation = doc.tables[3]
    untouched_text = {
        i: "\n".join(cell.text for row in table.rows for cell in row.cells)
        for i, table in enumerate(doc.tables)
        if i not in (2, 3)
    }

    # Insert the missing theoretical/model basis before the existing methods.
    scheme_anchor = next(row for row in scheme.rows if row.cells[0].text.strip() == "2. 拟采取的研究方法、技术路线及可行性分析")
    scheme_body_template = next(row for row in scheme.rows if row.cells[0].text.strip().startswith("（1）实验研究、公平对照与消融分析"))
    insert_rows_after(scheme, scheme_anchor, scheme_body_template, THEORY_ROWS)

    clear_and_set(find_cell_paragraph(scheme, "2.1 研究方法"), "2.2 研究方法", heading=True)
    clear_and_set(find_cell_paragraph(scheme, "2.2 技术路线"), "2.3 技术路线", heading=True)
    clear_and_set(find_cell_paragraph(scheme, "2.3 可行性分析"), "2.4 可行性分析", heading=True)

    # Make the feasibility paragraph explicitly connect the model to measurable evidence.
    feasibility = find_cell_paragraph_prefix(scheme, "在理论与方法方面")
    clear_and_set(
        feasibility,
        "在理论与方法方面，上述模型把研究对象明确分为基线B、候选算子O、晋级门控和带拒绝选项的路由器g，并把系统收益写成专项增益、非专项回归、超时和成本共同约束的目标。稳定失败门槛、候选质量函数、隐藏晋级条件和路由风险阈值均可在实验前固定；每个研究假设都有对应的对照、指标和失败判据。因此，本研究不是仅凭经验不断增加工具，而是能够通过逐题配对、隐藏验证和消融实验检验各模块贡献的可证伪方案。",
        heading=False,
    )

    # Turn the buried joint-result paragraph into a visible preliminary-results section.
    clear_and_set(find_cell_paragraph(foundation, "1.4 算子库、上下文路由与端到端实验"), "1.4 算子库与上下文路由原型", heading=True)
    buried_joint = find_cell_paragraph_prefix(foundation, "在全新种子的双算子联合隐藏实验中")
    result_anchor_row = next(
        row
        for row in foundation.rows
        if row.cells[0].text.strip().startswith("在全新种子的双算子联合隐藏实验中")
    )
    foundation_body_template = result_anchor_row
    clear_and_set(buried_joint, "1.5 初步实验结果", heading=True)
    insert_rows_after(
        foundation,
        result_anchor_row,
        foundation_body_template,
        [(False, text) for text in EXPERIMENT_ROWS],
    )
    clear_and_set(find_cell_paragraph(foundation, "1.5 可复现研究资产"), "1.6 可复现研究资产", heading=True)

    # Guard against accidental edits outside the requested sections.
    for i, before in untouched_text.items():
        after = "\n".join(cell.text for row in doc.tables[i].rows for cell in row.cells)
        if before != after:
            raise AssertionError(f"unexpected text change in table {i}")

    scheme_text = "\n".join(row.cells[0].text for row in scheme.rows)
    foundation_text = "\n".join(row.cells[0].text for row in foundation.rows)
    required_scheme = [
        "2.1 模型构建的理论依据与形式化定义",
        "F(x)=ok(x)",
        "2.2 研究方法",
        "2.3 技术路线",
        "2.4 可行性分析",
    ]
    required_results = [
        "1.5 初步实验结果",
        "91/132（68.9%）",
        "97/132（73.5%）",
        "96/144（66.7%）",
        "1.6 可复现研究资产",
    ]
    for token in required_scheme:
        if token not in scheme_text:
            raise AssertionError(f"missing scheme token: {token}")
    for token in required_results:
        if token not in foundation_text:
            raise AssertionError(f"missing result token: {token}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"scheme_rows={len(scheme.rows)}, foundation_rows={len(foundation.rows)}")


if __name__ == "__main__":
    main()
