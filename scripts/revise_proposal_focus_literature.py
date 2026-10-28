from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


SOURCE = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤 (1).docx')
OUTPUT = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_聚焦修订与文献扩充版.docx')

OLD_TITLE = '面向大语言模型智能体工作流的自动算子生成、可信验证与上下文路由研究'
NEW_TITLE = '面向大语言模型智能体工作流的可信算子自动生成与路由方法研究'
NEW_TITLE_EN = 'Research on Trustworthy Automatic Operator Generation and Routing Methods for Large Language Model Agent Workflows'


def set_paragraph_text(paragraph, text):
    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    p = paragraph._p
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return run


def set_cell_text(cell, text):
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        paragraph = cell.add_paragraph()
    else:
        paragraph = paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs)[1:]:
        cell._tc.remove(extra._p)
    for tbl in list(cell.tables):
        cell._tc.remove(tbl._tbl)
    return paragraph


def rebuild_single_column_table(table, items, template_indices):
    templates = {key: deepcopy(table.rows[idx]._tr) for key, idx in template_indices.items()}
    for tr in list(table._tbl.tr_lst):
        table._tbl.remove(tr)
    for kind, text in items:
        table._tbl.append(deepcopy(templates[kind]))
        set_cell_text(table.rows[-1].cells[0], text)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def put_results_table(cell):
    # Remove all existing body content and rebuild a compact comparison table.
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)
    table = cell.add_table(rows=4, cols=5)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Mm(25), Mm(27), Mm(27), Mm(47), Mm(35)]
    headers = ['确认实验', '原始AFlow', '路由系统', '专项任务', '路由与超时']
    data = [
        ['括号化算子', '91/132\n(68.9%)', '117/132\n(88.6%)', '8/36→36/36', '误路由0\n29→3'],
        ['受限数字素数算子', '97/132\n(73.5%)', '129/132\n(97.7%)', '4/36→36/36', '误路由0\n32→0'],
        ['双算子联合', '96/144\n(66.7%)', '132/144\n(91.7%)', '9/24→24/24；\n0/24→24/24', '误路由/抢路由0\n39→1'],
    ]
    rows = [headers] + data
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            c = table.cell(ri, ci)
            c.width = widths[ci]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            run.bold = ri == 0
            run.font.size = Pt(8.5 if ri else 9)
            run.font.name = '宋体'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    set_repeat_table_header(table.rows[0])
    # Word requires a trailing paragraph after a nested table.
    tail = cell.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)


doc = Document(SOURCE)

# Cover title.
for paragraph in doc.paragraphs:
    if OLD_TITLE in paragraph.text:
        set_paragraph_text(paragraph, paragraph.text.replace(OLD_TITLE, NEW_TITLE))

# Abstract and title metadata.
t0 = doc.tables[0]
set_cell_text(t0.rows[0].cells[-1], NEW_TITLE)
set_cell_text(t0.rows[1].cells[-1], NEW_TITLE_EN)
abstract = (
    '大语言模型智能体通过组合模型调用、工具与控制流程完成复杂任务。AFlow能够在人工预定义的基础算子集合上自动搜索算子组合与执行流程，'
    '但当既有算子无法覆盖新的稳定失败机制时，仅调整工作流结构难以补足相应能力。核心研究问题在于：如何从工作流运行失败中自动识别可复用的能力缺口，'
    '并将其转化为经过独立验证、能够安全接入工作流的新算子。围绕这一问题，本研究拟提出面向智能体工作流的可信算子自动生成与路由方法。首先，'
    '在排除评分异常、偶发超时和随机波动后判定稳定失败机制，并将能力缺口表示为包含输入输出、适用范围、算法约束和拒绝边界的结构化规格，'
    '利用受约束模板自动合成多个可复用候选算子；其次，以设计集、隐藏门控集和回归集相互隔离的方式评价候选的专项增益、稳定性与副作用，'
    '形成可追溯的晋级和版本管理机制；最后，依据问题特征、算子适用条件和历史证据进行选择性路由，在低置信度、输入越界或执行失败时回退至原始AFlow。'
    '前期已在MATH参数化问题族上完成两类自动生成专项算子及双算子联合验证：在全新隐藏种子实验中，系统整体正确率由66.7%提高至91.7%，两类专项任务均达到100%。'
    '结果初步说明所提方法具备端到端接入和多算子协同的可行性；自然题目分布下的整体收益、算子规模扩展和跨领域泛化仍需进一步验证。'
)
set_cell_text(t0.rows[3].cells[-1], abstract)
set_cell_text(t0.rows[4].cells[-1], '大语言模型智能体  智能体工作流  可信算子自动生成  选择性路由  算子库')
set_cell_text(t0.rows[5].cells[-1], 'large language model agents / agentic workflows / trustworthy automatic operator generation / selective routing / operator library')


# Rebuild the literature review/grounds section with a focused, denser literature chain.
t1 = doc.tables[1]
lit_items = [
    ('main', '1. 研究意义（工程价值）'),
    ('body', '大语言模型在知识问答、代码生成、数学推理和复杂决策等任务上表现出较强的语言理解与生成能力，但单次模型调用仍容易出现事实幻觉、精确计算错误、条件遗漏和长程执行失败。为提高复杂任务求解能力，研究者开始将模型调用、提示、记忆、外部工具、结果检查和控制流程组合为智能体工作流。工作流能够把开放任务分解为若干可执行步骤，并利用工具与环境反馈弥补语言模型在知识更新、精确计算和实际操作方面的不足。'),
    ('body', '为避免术语歧义，本文将“算子”定义为：可在智能体工作流中被独立调用、具有显式输入输出契约并完成一类可复用子任务的能力单元。其实现可以封装提示模板、模型调用、确定性程序或外部工具，同时附带适用条件、参数范围、拒绝与回退规则以及版本证据。函数或API是算子的执行载体，“技能”是更宽泛的能力封装，“节点”则是算子在具体工作流图中的一次实例。AFlow的基础算子由人工预定义，而本文研究的是从稳定失败机制出发自动形成、验证并接入工作流的新增算子。'),
    ('body', '现有自动工作流方法已经能够减少人工编排成本。以AFlow为代表的方法在人工预定义的基础算子集合上搜索算子组合、连接关系和执行流程。然而，当算子集合本身缺少完成某类任务所需的确定性能力时，仅改变组合顺序和控制逻辑不能消除这一能力缺口。因此，本研究把自动化对象由“预定义算子上的流程搜索”收束并推进到“可信算子自动生成”：验证用于回答生成算子是否具有可复用能力、能否晋级，路由用于回答晋级算子何时应当被调用，二者均服务于自动生成算子的可靠接入。'),
    ('body', '在工程应用方面，经过验证的参数化算子能够把重复、精确且可形式化的任务交给可复用程序，降低单题临时代码生成的语法错误、执行超时和重复成本；带有适用范围、失败边界和历史证据的算子库，可以为运行时选择提供比名称或简短描述更可靠的依据；选择性路由与回退机制则能够控制专项能力对普通任务的副作用。该研究有望为数学推理、代码生成和结构化数据处理等具有客观判定器的任务提供可扩展的能力补全方法。'),
    ('section', '2. 国内外研究现状（工程应用现状）'),
    ('sub', '2.1 推理增强、工具调用与反馈学习'),
    ('body', '大语言模型智能体研究首先从显式推理和工具增强展开。Chain-of-Thought通过中间推理步骤改善多步推理[1]；ReAct把推理、行动和环境观察组织为循环[2]；PAL将自然语言问题转换为程序并把计算交给解释器[3]；Toolformer研究模型自主决定何时调用工具及如何利用返回结果[4]。这些工作证明了外部执行模块能够补足模型的精确计算和环境交互能力，但工具集合、接口与适用范围通常仍由人工提供。'),
    ('body', '在执行反馈和能力积累方面，Reflexion利用语言反馈与情景记忆改进后续尝试[5]，Voyager持续生成可执行程序并把成功程序保存为技能以供检索和复用[6]。这类工作建立了“执行—反馈—积累”的基本范式，但通常从成功轨迹提取经验，较少研究如何从基线的重复失败中确认能力缺口，以及如何通过独立隐藏验证判定新能力是否可以安全进入工作流。'),
    ('sub', '2.2 提示、语言模型程序与智能体工作流自动优化'),
    ('body', '随着智能体流程复杂度提高，研究重点由人工调试转向自动优化。DSPy把语言模型应用表示为声明式模块与计算图，并依据任务指标优化提示和示例[7]；Promptbreeder通过进化机制生成和筛选提示[8]；TextGrad利用自然语言反馈修改提示、代码及其他文本变量[9]；Self-Refine通过模型自反馈迭代改进输出[10]。这些方法提高了语言模型程序的自动化程度，但多数是在既定模块边界内优化文本变量或单次输出。'),
    ('body', '自动智能体设计进一步扩大了搜索空间。Automated Design of Agentic Systems由元智能体以代码形式提出新的智能体设计[11]；AutoFlow以自然语言程序表示并优化工作流[12]；GPTSwarm将智能体表示为可优化计算图，同时调整节点提示和图连接[13]；AFlow通过蒙特卡洛树搜索、执行反馈和树状经验，在人工预定义算子集合上优化节点组合与控制流程[14]。这些工作解决了“如何组合已有能力”，但当基础能力模块缺失时仍需要扩展算子集合。'),
    ('sub', '2.3 工具创建、算子抽象与可复用能力生成'),
    ('body', '工具创建研究开始让模型从工具使用者转变为工具构造者。LATM把工具制作与工具使用分离，由模型生成可缓存的Python工具以服务后续同类任务[15]；CREATOR将抽象工具创建与具体任务执行解耦，并在数学与表格推理上验证了自动创建工具的价值[16]；CRAFT从训练样例生成代码解法，经过验证、抽象和去重形成可检索的专项工具集[17]。这些工作说明，参数化代码能够把多道题共享的求解机制沉淀为可复用能力，也为“设计与实现分离”和“先验证再复用”提供了直接依据。'),
    ('body', 'A²Flow进一步针对AFlow依赖人工预定义算子的问题，从专家案例生成初始算子，经聚类和抽象形成可复用执行算子[18]。与上述研究相比，本课题不以已有成功示例的抽象作为唯一入口，而是从基线的稳定失败机制出发，把机制判定、受约束程序合成、隐藏门控和失败边界共同纳入算子生成过程。关键差异不在于是否生成一段代码，而在于能否证明该代码代表可复用机制能力，并在未见数据上保持可靠。'),
    ('sub', '2.4 工具选择、技能路由与智能体评测'),
    ('body', '随着工具和技能数量增长，能力选择逐渐成为检索、排序和决策问题。SkillRouter在大规模候选技能中采用检索与重排序，并表明技能正文比名称和简短描述包含更充分的选择信号[19]；ToolLLM与ToolBench围绕大规模真实API构建工具学习数据和评价框架[20]；Gorilla通过检索感知训练提高API选择和调用可靠性[21]；HuggingGPT让语言模型依据功能描述选择并协调不同AI模型[22]。这些研究主要在既有工具集合中选择合适工具，尚未回答新生成算子应满足何种证据条件才能进入工具集合。'),
    ('body', 'AutoGen、MetaGPT和ChatDev分别从多智能体对话、软件工程标准流程和角色化协作角度验证了复杂工作流的工程可行性[23-25]。AgentBench和AgentBoard进一步提供多环境、长程交互和细粒度过程评价[26-27]。这些工作说明，仅报告最终答案正确率不足以解释智能体系统的行为；对于动态扩展的算子库，还需要记录路由决定、完整轨迹、执行时间、失败类型和版本变化。'),
    ('sub', '2.5 程序合成、选择性决策与可信评价基础'),
    ('body', '从方法基础看，程序合成把任务描述为“在候选程序空间中寻找满足规格的实现”[28]，语法引导合成进一步利用语法或模板约束候选空间[29]，为本研究的结构化算子规格与受约束模板提供理论依据。与针对单题自由生成代码不同，本文把稳定失败机制表示为可复用规格，再在受控假设空间中生成多个候选，以提高可执行性、可比较性和跨题复用能力。'),
    ('body', '从运行时决策看，SelectiveNet等带拒绝选项的选择性预测方法以覆盖率换取较低风险[30]；FrugalGPT和RouteLLM根据输入在不同模型之间动态分配，以平衡正确性、成本和延迟[31-32]。这些研究为算子路由提供了风险—覆盖视角，但模型路由通常把候选模型视为稳定黑盒，而自动生成算子的适用范围、失败边界和版本可靠性可能持续变化，因此还需把历史验证证据纳入路由。'),
    ('sub', '2.6 现有研究不足与本课题切入点'),
    ('body', '综合现有研究，仍存在三项相互关联的空白：第一，缺少从随机失败中识别可复用机制能力缺口，并把该缺口转化为受约束可执行算子的系统方法；第二，缺少在生成器不可见的隐藏数据上验证专项收益、回归、稳定性与失败边界的统一晋级机制；第三，缺少在算子持续增长时联合历史证据与风险进行选择、拒绝和回退的运行时方法。MATH提供可精确判定的数学问题场景[33]，SWE-bench等代码任务则提供以测试执行为判定依据的跨领域场景[34]，适合用于检验统一流程而非迁移具体专项代码。'),
    ('body', '据此，本课题以AFlow为通用工作流基线，把研究主线收束为可信算子自动生成：首先研究稳定失败机制的判定与机制约束下的算子程序合成；其次研究数据隔离条件下的可信晋级与生命周期管理；最后研究证据与风险联合驱动的选择性路由，使通过验证的算子能够安全接入原有工作流。验证与路由不是与“算子生成”并列的独立课题，而是自动生成算子从候选到可用能力所必需的质量保障和运行机制。'),
    ('section', '3. 主要参考文献及出处'),
]

refs = [
    '[1] WEI J, WANG X, SCHUURMANS D, et al. Chain-of-thought prompting elicits reasoning in large language models[C]//Advances in Neural Information Processing Systems. 2022, 35: 24824-24837.',
    '[2] YAO S, ZHAO J, YU D, et al. ReAct: Synergizing reasoning and acting in language models[C]//International Conference on Learning Representations. 2023.',
    '[3] GAO L, MIALON G, CHEN X, et al. PAL: Program-aided language models[C]//International Conference on Machine Learning. 2023: 10764-10799.',
    '[4] SCHICK T, DWIVEDI-YU J, DESSÌ R, et al. Toolformer: Language models can teach themselves to use tools[C]//Advances in Neural Information Processing Systems. 2023, 36: 68539-68551.',
    '[5] SHINN N, CASSANO F, GOPINATH A, et al. Reflexion: Language agents with verbal reinforcement learning[C]//Advances in Neural Information Processing Systems. 2023, 36: 8634-8652.',
    '[6] WANG G, XIE Y, JIANG Y, et al. Voyager: An open-ended embodied agent with large language models[J]. Transactions on Machine Learning Research, 2024.',
    '[7] KHATTAB O, SINGHVI A, MAHDIVAR A, et al. DSPy: Compiling declarative language model calls into self-improving pipelines[C]//International Conference on Learning Representations. 2024.',
    '[8] FERNANDO C, BANARSE D, MICHALEWSKI H, et al. Promptbreeder: Self-referential self-improvement via prompt evolution[EB/OL]. https://arxiv.org/abs/2309.16797, 2023-09-28/2026-08-27.',
    '[9] YUKSEKGONUL M, BIANCHI F, BOEN J, et al. TextGrad: Automatic “differentiation” via text[EB/OL]. https://arxiv.org/abs/2406.07496, 2024-06-11/2026-08-27.',
    '[10] MADAAN A, TANDON N, GUPTA P, et al. Self-Refine: Iterative refinement with self-feedback[C]//Advances in Neural Information Processing Systems. 2023, 36: 46534-46594.',
    '[11] HU S, LU C, CLUNE J. Automated design of agentic systems[C]//International Conference on Learning Representations. 2025.',
    '[12] LI Z, XU S, MEI K, et al. AutoFlow: Automated workflow generation for large language model agents[EB/OL]. https://arxiv.org/abs/2407.12821, 2024-07-01/2026-08-27.',
    '[13] ZHUGE M, WANG W, KIRSCH L, et al. GPTSwarm: Language agents as optimizable graphs[C]//International Conference on Machine Learning. 2024: 62743-62767.',
    '[14] ZHANG J, XIANG J, YU Z, et al. AFlow: Automating agentic workflow generation[C]//International Conference on Learning Representations. 2025.',
    '[15] CAI T, WANG X, MA T, et al. Large language models as tool makers[EB/OL]. https://arxiv.org/abs/2305.17126, 2023-05-26/2026-08-27.',
    '[16] QIAN C, HAN C, FUNG Y, et al. CREATOR: Tool creation for disentangling abstract and concrete reasoning of large language models[C]//Findings of the Association for Computational Linguistics: EMNLP 2023. Singapore: Association for Computational Linguistics, 2023: 6922-6939.',
    '[17] YUAN L, CHEN Y, WANG X, et al. CRAFT: Customizing LLMs by creating and retrieving from specialized toolsets[C]//International Conference on Learning Representations. 2024.',
    '[18] ZHAO M, WEI X, SHAO Y, et al. A²Flow: Automating agentic workflow generation via self-adaptive abstraction operators[EB/OL]. https://arxiv.org/abs/2511.20693, 2025-11-23/2026-08-27.',
    '[19] ZHENG Y, ZHANG Z, MA C, et al. SkillRouter: Retrieve-and-rerank skill selection for LLM agents at scale[EB/OL]. https://arxiv.org/abs/2603.22455, 2026-03-23/2026-08-27.',
    '[20] QIN Y, LIANG S, YE Y, et al. ToolLLM: Facilitating large language models to master 16000+ real-world APIs[C]//International Conference on Learning Representations. 2024.',
    '[21] PATIL S G, ZHANG T, WANG X, et al. Gorilla: Large language model connected with massive APIs[C]//Advances in Neural Information Processing Systems. 2024, 37.',
    '[22] SHEN Y, SONG K, TAN X, et al. HuggingGPT: Solving AI tasks with ChatGPT and its friends in Hugging Face[C]//Advances in Neural Information Processing Systems. 2023, 36.',
    '[23] WU Q, BANSAL G, ZHANG J, et al. AutoGen: Enabling next-gen LLM applications via multi-agent conversation[C]//Conference on Language Modeling. 2024.',
    '[24] HONG S, ZHUANG L, PAN J, et al. MetaGPT: Meta programming for a multi-agent collaborative framework[C]//International Conference on Learning Representations. 2024.',
    '[25] QIAN C, CONG X, YANG C, et al. ChatDev: Communicative agents for software development[C]//Proceedings of the 62nd Annual Meeting of the Association for Computational Linguistics. Bangkok: Association for Computational Linguistics, 2024: 15174-15186.',
    '[26] LIU X, YU H, ZHANG H, et al. AgentBench: Evaluating LLMs as agents[C]//International Conference on Learning Representations. 2024.',
    '[27] MA C, ZHANG J, ZHU Z, et al. AgentBoard: An analytical evaluation board of multi-turn LLM agents[C]//Advances in Neural Information Processing Systems. 2024, 37.',
    '[28] GULWANI S, POLOZOV O, SINGH R. Program synthesis[J]. Foundations and Trends in Programming Languages, 2017, 4(1-2): 1-119.',
    '[29] ALUR R, BODIK R, JUNIWAL G, et al. Syntax-guided synthesis[C]//Formal Methods in Computer-Aided Design. Portland: FMCAD Inc., 2013: 1-8.',
    '[30] GEIFMAN Y, EL-YANIV R. SelectiveNet: A deep neural network with an integrated reject option[C]//International Conference on Machine Learning. 2019: 2151-2160.',
    '[31] CHEN L, ZAHARIA M, ZOU J. FrugalGPT: How to use large language models while reducing cost and improving performance[J]. Transactions on Machine Learning Research, 2024.',
    '[32] ONG I, ALMAHAIRI A, WU V, et al. RouteLLM: Learning to route LLMs with preference data[C]//International Conference on Learning Representations. 2025.',
    '[33] HENDRYCKS D, BURNS C, KADAVATH S, et al. Measuring mathematical problem solving with the MATH dataset[C]//Advances in Neural Information Processing Systems: Datasets and Benchmarks Track. 2021, 1.',
    '[34] JIMENEZ C E, YANG J, WETTIG A, et al. SWE-bench: Can language models resolve real-world GitHub issues?[C]//International Conference on Learning Representations. 2024.',
]
lit_items.extend(('ref', ref) for ref in refs)
rebuild_single_column_table(t1, lit_items, {'main': 0, 'section': 4, 'sub': 4, 'body': 1, 'ref': 16})
for row in t1.rows[-len(refs):]:
    set_cant_split(row)


# Rebuild the research scheme around three aligned research questions.
t2 = doc.tables[2]
scheme_items = [
    ('main', '1. 研究目标、研究内容及拟解决的关键问题'),
    ('sub', '1.1 研究目标'),
    ('body', '本研究面向大语言模型智能体工作流在开放任务中出现的可复用能力缺口，围绕“可信算子自动生成”这一核心对象，研究从失败证据到可执行算子、从候选算子到可信能力、再从可信能力到安全调用的完整过程。总体目标是形成一套能够自动发现能力缺口、生成可复用算子并可靠接入AFlow工作流的方法。具体目标如下。'),
    ('body', '（1）实现失败机制驱动的算子自动生成。建立稳定失败机制的判定与表示方法，将重复能力缺口转化为结构化算子规格，并自动形成具有明确输入输出、适用范围和算法约束的可执行候选。'),
    ('body', '（2）实现自动生成算子的可信验证与生命周期管理。建立与生成过程隔离的能力验证、回归约束和证据记录机制，使算子的晋级、冻结、更新和淘汰具有一致且可追溯的依据。'),
    ('body', '（3）实现面向开放算子库的上下文感知路由与系统集成。依据任务上下文、算子适用边界和历史证据进行动态选择，在不确定或执行失败时拒绝调用并回退，从而把专项能力收益转化为完整工作流的整体净收益。'),
    ('sub', '1.2 研究内容'),
    ('body', '三项研究内容围绕同一主线逐级衔接：研究内容1回答“如何生成具有机制针对性和复用性的算子”，研究内容2回答“如何证明自动生成算子值得信任”，研究内容3回答“如何在多个可信算子共存时安全使用”。'),
    ('body', '研究内容1：稳定失败机制判定与受约束算子程序合成方法。首先从AFlow真实运行轨迹中区分评分异常、偶发超时、随机错误和可重复机制缺口，通过参数化问题族、重复运行和错误聚类形成稳定失败机制；然后将机制表示为包含输入域、输出契约、前置条件、求解约束和复杂度边界的结构化规格，在受约束模板空间中自动生成多个参数化候选，并利用静态检查、可执行性检查和候选多样性选择减少自由代码生成的不可控性。该内容研究的核心不是构造更多数据，而是“如何把自然语言描述的能力缺口转化为可靠的可执行能力”。'),
    ('body', '研究内容2：数据隔离条件下的可信算子晋级与生命周期管理方法。将数据划分为仅供生成器使用的设计集、生成器不可见的隐藏门控集以及用于副作用检查的回归集；通过逐题配对、重复运行、正确时间内完成率和错误审计评价候选。算子只有满足预先固定的专项增益、稳定性和回归约束，才连同版本、哈希、适用范围、失败边界和历史证据进入算子库；未通过候选仅归档，评分器变更则使用已保存回答离线重评并覆盖旧结论。该内容研究“如何在避免开发数据泄漏的条件下证明自动生成算子具有可复用能力”。'),
    ('body', '研究内容3：证据与风险联合驱动的选择性算子路由方法。对任务语义、表达式或程序结构、参数范围和算子证据进行联合建模，先检索可能适用的算子，再根据预期收益和风险重排；只有满足适用条件且风险低于阈值时才调用，否则拒绝并回退至原始AFlow。通过专项题、普通题、近似反例和多算子重叠场景评价路由召回、误触发、跨算子冲突、非专项回归和超时，并研究算子数量增加时的冲突消解和性能变化。'),
    ('sub', '1.3 拟解决的关键问题'),
    ('body', '（1）如何从随机失败中识别具有可复用性的机制能力缺口？单次失败可能来自评分、格式、超时或采样波动，只有跨参数、跨样例且重复出现的错误结构才适合作为算子生成目标。需要建立稳定失败机制的可操作定义、统计判定与聚类表示。'),
    ('body', '（2）如何将自然语言描述的能力缺口自动转化为可靠的可执行算子？能力描述与程序实现之间存在语义鸿沟，完全自由生成容易产生接口错误、超时和样例过拟合，过强模板又会限制新能力。需要研究机制规格、受约束程序空间、多候选生成和复用性之间的关系。'),
    ('body', '（3）如何在开放算子库中实现兼顾收益与风险的动态选择？专项算子有效并不意味着完整系统必然获益；当算子数量增加、适用范围重叠或输入接近失败边界时，需要同时控制漏调用、误调用、跨算子冲突和非专项回归，并在不确定时可靠回退。'),
    ('main', '2. 拟采取的研究方法、技术路线及可行性分析'),
    ('sub', '2.1 模型构建的理论依据与形式化定义'),
    ('body', '（1）稳定失败判定与机制归纳依据。设基线工作流为B，同一候选机制的问题族为Xj，B在不同参数问题和独立运行上的错误视为伯努利观测。只有当经验错误率不低于阈值τe、其置信下界仍高于最低操作阈值，且至少m道不同问题在r次运行中重复失败时，才把该问题族判定为稳定失败机制Mj。该定义把单题偶发错误提升为可检验的问题族性质，为后续归纳可复用规格提供证据基础。'),
    ('body', '（2）受约束程序合成依据。程序合成的基本目标是在候选程序空间中寻找满足规格的实现[28]，语法引导合成通过语法或模板缩小搜索空间[29]。本研究将机制Mj表示为规格Sj=(I,O,P,A,C)，其中I、O为输入输出契约，P为适用前置条件，A为允许的算法骨架，C为复杂度与安全约束；算子生成器仅依据Sj和设计集Ddes产生候选集合Oj。由此，“算子设计—代码实现”被转化为受规格和模板共同约束的程序搜索，而不是对单题任意生成代码。'),
    ('body', '（3）独立验证与约束风险最小化依据。隐藏门控集Dgate和回归集Dreg不参与候选生成。候选的晋级证据可表示为V(o)=ΔAccgate+λ1·Stability−λ2·Regression−λ3·Timeout−λ4·Cost，并同时满足ΔAccgate≥τg、Regression≤τr、Stability≥τs且不存在硬回归。该模型把“是否可信”定义为隐藏收益与副作用约束的联合判定，生成质量与晋级资格分别评价，避免用开发集上的单次高分替代可信性。'),
    ('body', '（4）带拒绝选项的选择性决策依据。设算子库为O={o1,…,oK}，路由器g(x)输出0表示使用基线，输出k表示尝试算子ok。完整系统在g(x)=k、输入满足适用条件且执行成功时输出ok(x)，其他情况输出B(x)。选择性预测通过拒绝部分不确定输入来控制风险[30]，模型路由研究则表明可根据输入平衡质量、成本和延迟[31-32]；本研究进一步把算子失败边界与历史验证证据纳入路由，使覆盖率、收益和风险能够共同调节。'),
    ('body', '基于上述模型，论文检验三个可证伪假设：H1，受约束自动生成算子在生成器不可见的机制隐藏集上优于原始AFlow；H2，可信晋级能够筛除开发集高分但存在隐藏回归或不稳定性的候选；H3，证据感知的选择性路由能够把专项增益转化为混合分布下的整体净收益，并在算子数量增加时把误触发和跨算子冲突控制在预设范围。'),
    ('sub', '2.2 研究方法'),
    ('body', '（1）实验研究与公平对照。统一数据、执行模型、运行次数、超时和工具权限，比较原始AFlow、AFlow加通用单题代码生成、人工实现专项算子、自动生成专项算子、自动生成算子加证据感知路由，以及理想路由或人工规则上界。逐层比较用于区分代码执行、人工知识、自动生成和路由各自贡献。'),
    ('body', '（2）冻结盲测与消融分析。原119题仅作为开发诊断和回归材料；正式结论使用预先固定协议、生成器不可见且仅确认一次的全新隐藏集。分别消融稳定失败筛选、结构化规格、受约束模板、多候选选择、隐藏晋级、拒绝回退和历史证据，分析各机制对结果的作用。'),
    ('body', '（3）定量评价与逐题审计。以correct_within_timeout为主指标，同时报告整体与专项正确率、逐题胜/平/负、非专项回归、路由召回、误触发、跨算子路由、超时、执行成本和重复运行方差；利用保存的完整回答区分求解、评分、路由和工具执行错误。'),
    ('body', '（4）案例研究与跨领域复现。分析成功和失败算子的规格、实现与边界，解释系统为何生成、晋级、调用或拒绝。跨领域实验复用机制表示、候选生成、可信门控和路由框架，并为代码任务或结构化数据任务配置独立可执行判定器，以检验方法通用性。'),
    ('sub', '2.3 技术路线'),
    ('body', '第一阶段：失败机制驱动的算子自动生成。运行原始AFlow并保存轨迹，完成评分审计与错误聚类；通过baseline-only参数化校准确认稳定能力缺口；将入选机制转化为结构化规格，使用受约束模板自动生成多个可复用候选，并完成接口、静态和冒烟检查。'),
    ('body', '第二阶段：可信晋级与证据化管理。使用相互隔离的设计集、隐藏门控集和回归集开展候选锦标赛与独立确认；保存完整回答、耗时、错误类型、版本与文件哈希。达到预设门槛的算子进入版本化算子库，未通过候选归档，旧隐藏种子冻结。'),
    ('body', '第三阶段：选择性路由与系统验证。依据任务特征、适用条件、失败边界和历史证据检索、重排并选择算子；在专项题、普通题、近似反例和重叠场景组成的全新混合隐藏集上开展单算子、多算子和跨领域实验，验证整体收益、风险控制与扩展性。'),
    ('sub', '2.4 可行性分析'),
    ('body', '在理论与方法方面，稳定失败判定、受约束程序合成、约束风险最小化和带拒绝选项的选择性决策分别为三个研究内容提供依据。机制阈值、生成空间、晋级门槛和路由风险均可在实验前形式化并固定，每项研究内容都有对应对照、消融指标和失败判据，因此方案可以被实验验证或否定。'),
    ('body', '在工程与实验基础方面，现有系统已经实现运行日志保存、评分器离线重评、baseline-only筛选、参数化问题族生成、多候选锦标赛、算子版本管理和上下文路由。两个自动生成专项算子已分别通过全新隐藏集确认，双算子联合实验中整体正确率由66.7%提高至91.7%，且未发生普通题误路由或跨算子抢路由，说明完整流程已经具备可运行原型。详细协议和结果列于“四、研究条件与基础”。'),
    ('body', '在资源条件方面，研究不需要重新训练大语言模型，参数化数据生成和确定性算子执行成本可控；模型、随机种子、运行次数、超时和评分器版本均可固定，完整回答和哈希可长期保存。后续仍需补充自然分布、多算子规模扩展、不同执行模型和跨领域实验，以形成完整论文证据链。'),
    ('main', '3. 创新之处'),
    ('body', '（1）提出稳定失败机制约束下的算子自动生成方法。将随机运行错误归纳为跨参数、跨样例重复出现的机制能力缺口，并表示为包含输入输出、前置条件、算法骨架和复杂度边界的算子规格；在规格与模板限定的程序空间中自动生成多个参数化候选。该创新聚焦“如何生成具有机制针对性、可执行性和复用性的算子”，不把隐藏验证结果混入生成目标。'),
    ('body', '（2）提出面向自动生成算子的独立证据晋级方法。通过设计集、隐藏门控集与回归集隔离，把专项增益、稳定性、回归、超时和成本转化为预先固定的晋级约束；结合完整回答、版本哈希、离线重评和种子冻结管理算子生命周期。该创新聚焦“如何证明自动生成算子值得信任”，使入库资格由独立证据而非开发样例得分决定。'),
    ('body', '（3）提出证据与风险联合驱动的选择性算子路由方法。路由器联合任务语义与结构、算子适用范围、失败边界和历史可靠性决定调用、拒绝或回退，并以混合隐藏集中的系统净收益、误触发、跨算子冲突和非专项回归评价路由。该创新聚焦“如何安全使用已经通过验证的算子”，使开放算子库在扩展时仍能控制运行风险。'),
]
rebuild_single_column_table(t2, scheme_items, {'main': 0, 'sub': 1, 'body': 2})


# Turn the preliminary experiments in the research foundation into an explicit table.
t3 = doc.tables[3]
set_cell_text(t3.rows[6].cells[0], '1.3 失败机制驱动的算子自动生成原型')
set_cell_text(t3.rows[8].cells[0], '1.4 可信验证、算子库与上下文路由原型')
set_cell_text(t3.rows[9].cells[0], '目前已形成从候选生成到可信接入的可运行原型。候选算子在设计集上完成结构化规格生成、受约束实现、静态检查和冒烟测试，随后进入与生成阶段隔离的隐藏门控和回归审计；通过后记录名称、版本、适用条件、输入范围、拒绝与回退规则、验证结果和文件哈希。路由器根据自然语言意图、结构特征和参数范围选择算子，并在不适用或执行失败时回到原始AFlow。现已完成ReusableParenthesizationIntervalDPTool v2和ReusableRestrictedDigitPrimeEnumerator v2两个自动生成参数化专项算子。')
set_cell_text(t3.rows[10].cells[0], '1.5 初步实验结果')
set_cell_text(t3.rows[11].cells[0], '实验协议。三组确认实验均使用候选生成器未见过的全新随机种子隐藏集，执行模型为deepseek-v4-flash，每道题独立运行3次，单次超时180秒，主指标为correct_within_timeout。实验前冻结算子版本、路由规则、成功门槛和文件哈希；确认过程中不修改候选，使用math_equivalence_v7评分器，并保存逐题回答和路由记录。')
set_cell_text(t3.rows[12].cells[0], '表1  自动生成专项算子接入AFlow的初步端到端结果')
put_results_table(t3.rows[13].cells[0])
set_cell_text(t3.rows[14].cells[0], '结果分析。两个单算子实验均把专项题正确率提高到100%，且普通题与近似反例误路由均为0；双算子联合时，整体正确率由96/144提高到132/144，两类专项任务共48次观测全部正确，未发生普通题误路由或跨算子抢路由。联合实验出现3次非专项新增错误，经逐题日志审计均未触发专项路由，因此不能归因于两个算子。')
set_cell_text(t3.rows[15].cells[0], '阶段性结论。结果支持“通过独立门控的自动生成算子能够接入AFlow，并在参数化混合隐藏集上产生端到端净收益”，也初步说明两个算子共同存在时可以协同工作。但该结果仍属于MATH参数化问题族的可行性证据，不能替代自然题目分布、更多失败机制、不同执行模型和跨领域实验；后续将按照研究方案中的公平对照与冻结协议继续验证。')


# Align expected outcomes with the three-part research structure.
t4 = doc.tables[4]
set_cell_text(t4.rows[11].cells[0], '（1）形成一套面向稳定失败机制的可信算子自动生成方法，给出机制判定、结构化规格、受约束程序合成和多候选生成的形式化定义与实现。 （2）形成自动生成算子的可信晋级、证据化生命周期管理和选择性路由方法，完成公平对照、主要模块消融、多算子联合及不同执行模型实验。 （3）形成可复现的原型系统、版本化算子库和实验资产，在MATH之外至少选择一种具有客观判定标准的任务复现统一流程，完成学位论文及相关学术成果。')


# Replace any remaining old title occurrence in tables without disturbing other text.
for table in doc.tables:
    for row in table.rows:
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            for paragraph in cell.paragraphs:
                if OLD_TITLE in paragraph.text:
                    set_paragraph_text(paragraph, paragraph.text.replace(OLD_TITLE, NEW_TITLE))

doc.core_properties.title = NEW_TITLE
doc.save(OUTPUT)
print(OUTPUT)
