from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_聚焦修订_含5至6月实验版.docx')
OUTPUT = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_聚焦修订_含探索性实验版.docx')


REPLACEMENTS = {
    '2026年5—6月开发性探索实验': '开发性探索实验',
    '2026年5月至6月围绕': '前期围绕',
}


def replace_paragraph_text(paragraph):
    old_text = paragraph.text
    new_text = old_text
    for old, new in REPLACEMENTS.items():
        new_text = new_text.replace(old, new)
    if new_text == old_text:
        return 0

    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    p = paragraph._p
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    run = paragraph.add_run(new_text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return 1


def walk_table(table):
    changed = 0
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            key = cell._tc
            if key in seen:
                continue
            seen.add(key)
            for paragraph in cell.paragraphs:
                changed += replace_paragraph_text(paragraph)
            for nested in cell.tables:
                changed += walk_table(nested)
    return changed


doc = Document(SOURCE)
changed = sum(replace_paragraph_text(paragraph) for paragraph in doc.paragraphs)
for table in doc.tables:
    changed += walk_table(table)

if changed != 3:
    raise RuntimeError(f'expected 3 paragraph replacements, found {changed}')

doc.core_properties.comments = '删除探索性实验标题、表题及说明中的具体月份表述。'
doc.save(OUTPUT)
print(OUTPUT)
