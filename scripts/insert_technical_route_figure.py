from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import _Cell


SOURCE = Path(
    "/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_技术路线与国内文献补充版.docx"
)
FIGURE = Path(
    "/Users/verneywang/Downloads/_cgi-bin_mmwebwx-bin_webwxgetmsgimg__&MsgID=4694923557350124158&skey=@crypt_64a17cb_1c9dd370cae36947686508bb4f218786&mmweb_appid=wx_webfilehelper.jpeg"
)
OUTPUT = Path(
    "/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_技术路线图已插入版.docx"
)

CAPTION = "图1  面向大语言模型智能体工作流的可信算子自动生成与路由技术路线"
ALT_TEXT = "技术路线图：稳定失败识别、受约束算子生成、可信验证晋级和上下文感知路由四层流程"


def set_east_asian_font(run, name: str, size_pt: float) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def clear_cell_preserving_properties(cell: _Cell) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.add_p()


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    doc = Document(SOURCE)
    target_table = None
    route_row_index = None
    for table in doc.tables:
        for index, row in enumerate(table.rows):
            text = " ".join(cell.text.strip() for cell in row.cells)
            if text == "2.3 技术路线":
                target_table = table
                route_row_index = index
                break
        if target_table is not None:
            break

    if target_table is None or route_row_index is None:
        raise RuntimeError("未找到‘2.3 技术路线’所在表格")

    next_heading_index = None
    for index in range(route_row_index + 1, len(target_table.rows)):
        text = " ".join(cell.text.strip() for cell in target_table.rows[index].cells)
        if text == "2.4 可行性分析":
            next_heading_index = index
            break
    if next_heading_index is None:
        raise RuntimeError("未找到‘2.4 可行性分析’，无法确定插图位置")

    # 复制相邻正文行以完整继承模板的表格宽度、边框和单元格属性，
    # 再把新行放到“2.4 可行性分析”之前。
    template_row = target_table.rows[next_heading_index - 1]
    reference_row = target_table.rows[next_heading_index]
    new_tr = deepcopy(template_row._tr)
    reference_row._tr.addprevious(new_tr)
    new_row = target_table.rows[next_heading_index]
    cell = new_row.cells[0]
    clear_cell_preserving_properties(cell)
    prevent_row_split(new_row)

    picture_paragraph = cell.paragraphs[0]
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.paragraph_format.keep_together = True
    picture_run = picture_paragraph.add_run()
    inline_shape = picture_run.add_picture(str(FIGURE), width=Mm(148))
    inline_shape._inline.docPr.set("title", "可信算子自动生成与路由技术路线")
    inline_shape._inline.docPr.set("descr", ALT_TEXT)

    caption_paragraph = cell.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_paragraph.paragraph_format.keep_together = True
    caption_run = caption_paragraph.add_run(CAPTION)
    set_east_asian_font(caption_run, "宋体", 9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
