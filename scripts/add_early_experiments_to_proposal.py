from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import _Cell


SOURCE = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_聚焦修订与文献扩充版.docx')
OUTPUT = Path('/Users/verneywang/Downloads/研究生学位（毕业）论文开题报告_马佳彤_聚焦修订_含5至6月实验版.docx')


def set_paragraph_text(paragraph, text):
    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    p = paragraph._p
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    return run


def set_cell_text(cell, text):
    paragraphs = list(cell.paragraphs)
    paragraph = paragraphs[0] if paragraphs else cell.add_paragraph()
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs)[1:]:
        cell._tc.remove(extra._p)
    for table in list(cell.tables):
        cell._tc.remove(table._tbl)
    return paragraph


def set_cell_margins(cell, top=65, start=80, bottom=65, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:tblHeader')) is None:
        node = OxmlElement('w:tblHeader')
        node.set(qn('w:val'), 'true')
        tr_pr.append(node)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:cantSplit')) is None:
        tr_pr.append(OxmlElement('w:cantSplit'))


def insert_cloned_row_before(table, reference_tr, template_tr, text=''):
    new_tr = deepcopy(template_tr)
    reference_tr.addprevious(new_tr)
    cell = _Cell(new_tr.tc_lst[0], table)
    set_cell_text(cell, text)
    return new_tr, cell


def put_exploratory_table(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'):
            cell._tc.remove(child)

    table = cell.add_table(rows=6, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Mm(30), Mm(48), Mm(50), Mm(33)]
    rows = [
        ['实验阶段', '主要开发结果', '对研究方案的作用', '证据定位'],
        [
            'MATH算子库扩展',
            '3算子最佳0.8319；6算子无验证时最佳0.8235，并出现4个零分轮；加入AST检查后最佳0.8403，零分轮降至1个。',
            '自由代码错误会掩盖算子效果，候选进入评测前需要静态与结构检查。',
            '119题开发集消融',
        ],
        [
            '冒烟与保留策略',
            'AST与冒烟测试结合后，11轮开发运行未出现零分；加入退火式最优保留时最佳0.7983，移除后恢复到0.8403。',
            '验证应快速识别失效候选，但过强的保留约束可能限制工作流探索。',
            '开发性观察',
        ],
        [
            'v4自动进化闭环',
            '首次跑通瓶颈检测、失败分析、候选生成、验证和安装；4次生成中2个候选存活，但均偏向重复验证类，最佳分数0.7899。',
            '证明自动闭环可运行，同时暴露生成方向单一和能力增益不足。',
            '原型可行性',
        ],
        [
            'v5—v6约束与反馈',
            'v5候选存活1/4、冒烟拒绝率37.5%；v6加入错误反馈与类名同步后存活4/4、拒绝率20%，但新算子未被优化器稳定采用。',
            '受约束生成和错误反馈改善可执行性；还需独立能力门控和上下文路由。',
            '开发性观察',
        ],
        [
            '跨任务辅助实验',
            'GSM8K高基线实验暴露评分与回归问题；HumanEval的33题子集上，AFlow工作流搜索最佳分数由0.8788提高到0.9697。',
            '说明实验基础设施可迁移，但尚不能证明自动算子生成的跨领域有效性。',
            '辅助证据',
        ],
    ]

    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            target = table.cell(ri, ci)
            target.width = widths[ci]
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(target)
            paragraph = target.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 3) or ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            run.bold = ri == 0
            run.font.size = Pt(8.2 if ri else 8.8)
            run.font.name = '宋体'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
        set_cant_split(table.rows[ri])
    set_repeat_header(table.rows[0])

    tail = cell.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)


doc = Document(SOURCE)
foundation = doc.tables[3]

# Preserve the original section and insert a clearly separated development-evidence layer.
set_cell_text(foundation.rows[10].cells[0], '1.5 前期实验结果')
reference_tr = foundation.rows[11]._tr
heading_template = foundation.rows[1]._tr
body_template = foundation.rows[2]._tr
caption_template = foundation.rows[12]._tr
table_template = foundation.rows[13]._tr

insert_cloned_row_before(
    foundation,
    reference_tr,
    heading_template,
    '1.5.1 2026年5—6月开发性探索实验',
)
insert_cloned_row_before(
    foundation,
    reference_tr,
    body_template,
    '实验定位。2026年5月至6月围绕固定算子空间扩展、候选代码验证和自动进化闭环开展了多轮探索。相关实验主要使用被反复分析的MATH 119题开发集，部分还使用GSM8K和HumanEval子集，未采用当前的冻结隐藏协议，因此用于说明研究问题的形成、原型可行性和方法选择依据，不作为最终有效性结论。',
)
insert_cloned_row_before(
    foundation,
    reference_tr,
    caption_template,
    '表1  2026年5—6月开发性探索实验及其证据定位',
)
_, exploratory_table_cell = insert_cloned_row_before(
    foundation,
    reference_tr,
    table_template,
    '',
)
put_exploratory_table(exploratory_table_cell)
insert_cloned_row_before(
    foundation,
    reference_tr,
    body_template,
    '探索性结论。早期实验没有证明“自动生成算子必然优于原始AFlow”，但形成了三项直接支撑本研究方案的认识：一是算子候选需要在进入正式评测前完成静态与运行检查；二是仅提高候选代码存活率仍不能保证能力收益，必须使用与生成隔离的隐藏门控和回归约束；三是新算子若不能被工作流正确选择，生成阶段的改进无法转化为端到端收益。这些发现分别推动了受约束算子生成、可信晋级和选择性路由三项研究内容。',
)
insert_cloned_row_before(
    foundation,
    reference_tr,
    heading_template,
    '1.5.2 冻结隐藏集上的端到端确认实验',
)

# The existing clean confirmation evidence follows the inserted exploratory section.
for row in foundation.rows:
    if row.cells[0].text.strip() == '表1  自动生成专项算子接入AFlow的初步端到端结果':
        set_cell_text(row.cells[0], '表2  自动生成专项算子接入AFlow的初步端到端结果')
        break

doc.core_properties.comments = '补充2026年5—6月开发性探索实验，并与冻结隐藏集确认结果分层呈现。'
doc.save(OUTPUT)
print(OUTPUT)
